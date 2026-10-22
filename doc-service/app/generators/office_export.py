import io
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from docx import Document
from fastapi.responses import StreamingResponse

from app.core.s3 import upload_buffer_to_s3
from app.schemas.documents import TableExportRequest

def generate_xlsx(req: TableExportRequest):
    """Generates a styled Excel (.xlsx) spreadsheet using openpyxl."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    # Title Row
    ws.append([req.title])
    ws["A1"].font = Font(size=14, bold=True, color="0F172A")
    if req.subtitle:
        ws.append([req.subtitle])
    ws.append([])

    # Header Row
    ws.append(req.headers)
    header_row_idx = ws.max_row
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx in range(1, len(req.headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data Rows
    for row in req.rows:
        ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"{req.title.replace(' ', '_')}.xlsx"
    if req.saveToS3:
        bucket = os.getenv("S3_BUCKET_DOCS", "printgoo-docs")
        key = req.s3Key or f"exports/{filename}"
        upload_buffer_to_s3(
            buf,
            bucket,
            key,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        return {"ok": True, "bucket": bucket, "s3Key": key, "filename": filename}

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )




def generate_docx(req: TableExportRequest):
    """Generates a Word (.docx) document with a formatted table using python-docx."""
    doc = Document()
    doc.add_heading(req.title, level=1)
    if req.subtitle:
        doc.add_paragraph(req.subtitle)

    if req.headers:
        table = doc.add_table(rows=1, cols=len(req.headers))
        table.style = "Table Grid"
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(req.headers):
            hdr_cells[i].text = str(header_text)

        for row in req.rows:
            row_cells = table.add_row().cells
            for i, val in enumerate(row):
                if i < len(row_cells):
                    row_cells[i].text = str(val)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    filename = f"{req.title.replace(' ', '_')}.docx"
    if req.saveToS3:
        bucket = os.getenv("S3_BUCKET_DOCS", "printgoo-docs")
        key = req.s3Key or f"exports/{filename}"
        upload_buffer_to_s3(
            buf,
            bucket,
            key,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        return {"ok": True, "bucket": bucket, "s3Key": key, "filename": filename}

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
